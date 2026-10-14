"""Reports and exports (docs/ARCHITECTURE.md §4 /console/reports/<name>/, §8 test_reports_exports).

Ageing and pending-by-approver reports show the right counts for invoices in known states; CSV and
Excel downloads (export-csv-btn / export-xlsx-btn) carry the agreed header columns (the column order
is a contract with Finance, see apps/reports/exports.py) and the same rows as the page.
"""
import csv
import io
from datetime import date, timedelta
from decimal import Decimal

import pytest
from openpyxl import load_workbook
from playwright.sync_api import expect

from .conftest import login

INVOICE_HEADERS = [
    "Portal ref", "Vendor code", "Vendor", "Supplier GSTIN", "Invoice no", "Invoice date", "FY", "Cost centre",
    "Taxable", "CGST", "SGST", "IGST", "Total", "TDS section", "TDS amount", "Status", "Due date", "Paid", "Outstanding",
]
HEADERS = {
    "ageing": ["Bucket", "Days"] + INVOICE_HEADERS,
    "pending-by-approver": ["Approver", "Email", "Level", "Pending count", "Pending amount", "Oldest since"],
    "vendor-master": ["Vendor code", "Legal name", "Trade name", "Category", "PAN (last 4)", "Primary GSTIN", "MSME status", "Udyam no", "Contact", "Email", "Status", "Activated"],
    "payments": ["Paid on", "Portal ref", "Vendor code", "Vendor", "Invoice no", "Amount", "TDS deducted", "UTR"],
}


# --- local helpers ---------------------------------------------------------------------------

def _submit(seed, *, total, invoice_no, days_old=0, vendor="vendor_a", by_email="vendora.user@test.local"):
    from apps.accounts.models import User
    from apps.core.seed import sample_invoice
    from apps.invoices import services

    v = seed[vendor]
    by = User.objects.get(email=by_email)
    inv = sample_invoice(v, by, total=total, invoice_no=invoice_no)
    data = {
        "supplier_gstin": v.primary_gstin.gstin, "invoice_no": invoice_no, "invoice_date": date.today() - timedelta(days=days_old),
        "taxable": inv.taxable, "cgst": inv.cgst, "sgst": inv.sgst, "igst": inv.igst, "total": inv.total,
        "cost_centre": seed["cost_centre"],
    }
    return services.submit(inv, by=by, data=data)


def _approve_to(seed, inv, target):
    """Approve level by level until status == target (a PENDING_* status or APPROVED)."""
    from apps.invoices import services

    actors = {"FIN_PRE": seed["finance"], "L1": seed["l1"], "L2": seed["l2"], "L3": seed["l3"]}
    while inv.status != target:
        fin = {"tds_section": "194J", "tds_amount": Decimal("0")} if inv.current_level == "FIN_PRE" else None
        inv = services.act(inv, by=actors[inv.current_level], action="APPROVE", fin_data=fin)
    return inv


@pytest.fixture
def known_invoices(seed):
    """Invoices in known states and ages.

    ageing (open = FIN_PRECHECK..PARTLY_PAID incl. QUERY_RAISED): 0-30 -> 4, 31-60 -> 2, 61-90 -> 1, 90+ -> 1
    pending by approver: finance 1, finance2 1, l1 2, l2 1, l3 1
    """
    from apps.core.seed import sample_invoice
    from apps.invoices import services

    s = lambda **kw: _submit(seed, **kw)  # noqa: E731
    out = {}
    out["a"] = s(total=50000, invoice_no="RPT-A", days_old=0)                                  # FIN_PRECHECK
    out["b"] = _approve_to(seed, s(total=50000, invoice_no="RPT-B", days_old=10), "PENDING_L1")
    out["c"] = _approve_to(seed, s(total=80000, invoice_no="RPT-C", days_old=20), "PENDING_L1")
    out["d"] = _approve_to(seed, s(total=500000, invoice_no="RPT-D", days_old=45), "PENDING_L2")
    out["e"] = _approve_to(seed, s(total=1500000, invoice_no="RPT-E", days_old=50), "PENDING_L3")
    out["f"] = _approve_to(seed, s(total=50000, invoice_no="RPT-F", days_old=75), "APPROVED")
    g = _approve_to(seed, s(total=50000, invoice_no="RPT-G", days_old=120), "APPROVED")
    services.record_payment(g, by=seed["finance"], amount=Decimal("10000"), paid_on=date.today(), utr="UTRRPTG01")
    out["g"] = g                                                                                # PARTLY_PAID
    h = _approve_to(seed, s(total=50000, invoice_no="RPT-H", days_old=5), "PENDING_L1")
    out["h"] = services.act(h, by=seed["l1"], action="QUERY", comment="Which PO?")              # QUERY_RAISED
    i = s(total=50000, invoice_no="RPT-I", days_old=100)
    out["i"] = services.act(i, by=seed["finance"], action="REJECT", comment="Duplicate")        # REJECTED (not open)
    j = _approve_to(seed, s(total=50000, invoice_no="RPT-J", days_old=40), "APPROVED")
    services.record_payment(j, by=seed["finance"], amount=Decimal("50000"), paid_on=date.today(), utr="UTRRPTJ01")
    out["j"] = j                                                                                # PAID (not open)
    sample_invoice(seed["vendor_a"], seed["finance"], total=1000, invoice_no="RPT-DRAFT")      # DRAFT (not open)
    return out


def _download(page, testid):
    with page.expect_download() as dl:
        page.get_by_test_id(testid).click()
    return dl.value


def _csv_rows(download):
    with open(download.path(), encoding="utf-8") as fh:
        return list(csv.reader(io.StringIO(fh.read())))


def _xlsx_rows(download):
    with open(download.path(), "rb") as fh:
        wb = load_workbook(io.BytesIO(fh.read()))
    return [list(r) for r in wb.active.iter_rows(values_only=True)]


# --- tests -----------------------------------------------------------------------------------

def test_ageing_report_counts_and_exports(page, seed, known_invoices):
    login(page, "finance@test.local")
    page.goto("/console/reports/ageing/")
    expected = {"0-30": 4, "31-60": 2, "61-90": 1, "90+": 1}
    for bucket, n in expected.items():
        expect(page.get_by_test_id(f"ageing-count-{bucket}")).to_have_text(str(n))
    rows = page.get_by_test_id("report-table").get_by_test_id("report-row")
    expect(rows).to_have_count(8)
    for key in "abcdefgh":
        expect(rows.filter(has_text=known_invoices[key].portal_ref)).to_have_count(1)
    for key in "ij":
        expect(rows.filter(has_text=known_invoices[key].portal_ref)).to_have_count(0)

    csv_rows = _csv_rows(_download(page, "export-csv-btn"))
    assert csv_rows[0] == HEADERS["ageing"]
    body = csv_rows[1:]
    assert len(body) == 8
    by_ref = {r[2]: r for r in body}
    assert set(by_ref) == {known_invoices[k].portal_ref for k in "abcdefgh"}
    assert by_ref[known_invoices["d"].portal_ref][0] == "31-60"
    assert by_ref[known_invoices["g"].portal_ref][0] == "90+"
    assert by_ref[known_invoices["g"].portal_ref][HEADERS["ageing"].index("Status")] == "PARTLY_PAID"
    assert Decimal(by_ref[known_invoices["g"].portal_ref][HEADERS["ageing"].index("Outstanding")]) == Decimal("40000")
    buckets = [r[0] for r in body]
    assert {b: buckets.count(b) for b in expected} == expected

    x = _xlsx_rows(_download(page, "export-xlsx-btn"))
    assert x[0] == HEADERS["ageing"]
    assert len(x) == 9
    assert sorted(r[0] for r in x[1:]) == sorted(buckets)
    assert {r[2] for r in x[1:]} == set(by_ref)


def test_pending_by_approver_counts_and_exports(page, seed, known_invoices):
    login(page, "auditor@test.local")
    page.goto("/console/reports/pending-by-approver/")
    expected = {"finance@test.local": 1, "finance2@test.local": 1, "l1@test.local": 2, "l2@test.local": 1, "l3@test.local": 1}
    for email, n in expected.items():
        expect(page.get_by_test_id(f"pending-count-{email}")).to_have_text(str(n))
    expect(page.get_by_test_id("report-table").get_by_test_id("report-row")).to_have_count(len(expected))
    for absent in ("vm@test.local", "admin@test.local", "auditor@test.local"):
        expect(page.get_by_test_id(f"pending-count-{absent}")).to_have_count(0)

    csv_rows = _csv_rows(_download(page, "export-csv-btn"))
    assert csv_rows[0] == HEADERS["pending-by-approver"]
    got = {r[1]: r for r in csv_rows[1:]}
    assert {e: int(r[3]) for e, r in got.items()} == expected
    assert got["l1@test.local"][2] == "L1 approval"
    assert Decimal(got["l1@test.local"][4]) == Decimal("130000")
    assert Decimal(got["l3@test.local"][4]) == Decimal("1500000")

    x = _xlsx_rows(_download(page, "export-xlsx-btn"))
    assert x[0] == HEADERS["pending-by-approver"]
    assert {r[1]: int(r[3]) for r in x[1:]} == expected


@pytest.mark.parametrize("name", ["vendor-master", "payments", "ageing", "pending-by-approver"])
def test_report_export_headers(page, seed, known_invoices, name):
    login(page, "finance@test.local")
    page.goto(f"/console/reports/{name}/")
    n_rows = page.get_by_test_id("report-table").get_by_test_id("report-row").count()

    csv_dl = _download(page, "export-csv-btn")
    assert csv_dl.suggested_filename == f"{name}.csv"
    rows = _csv_rows(csv_dl)
    assert rows[0] == HEADERS[name]
    assert len(rows) - 1 == n_rows

    xlsx_dl = _download(page, "export-xlsx-btn")
    assert xlsx_dl.suggested_filename == f"{name}.xlsx"
    x = _xlsx_rows(xlsx_dl)
    assert x[0] == HEADERS[name]
    assert len(x) - 1 == n_rows
    if name == "payments":
        assert sorted(r[HEADERS[name].index("UTR")] for r in x[1:]) == ["UTRRPTG01", "UTRRPTJ01"]
    if name == "vendor-master":
        assert sorted(r[0] for r in x[1:]) == ["V00001", "V00002"]


def test_console_invoice_list_export_headers(page, seed, known_invoices):
    login(page, "finance@test.local")
    page.goto("/console/invoices/?status=PENDING_L1")
    rows = _csv_rows(_download(page, "export-csv-btn"))
    assert rows[0] == INVOICE_HEADERS
    assert sorted(r[0] for r in rows[1:]) == sorted([known_invoices["b"].portal_ref, known_invoices["c"].portal_ref])
    x = _xlsx_rows(_download(page, "export-xlsx-btn"))
    assert x[0] == INVOICE_HEADERS
    assert len(x) == 3
