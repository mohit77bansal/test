"""Request-level tests: auth, role landing, vendor isolation (404s), exports, page rendering."""
import csv
import io
import re
from decimal import Decimal

import pytest
from django.core import mail
from django.urls import reverse
from openpyxl import load_workbook

from apps.invoices import services
from tests.conftest import fake_pdf

PW = "Passw0rd!234"


def login(client, email, password=PW):
    resp = client.post("/login/", {"email": email, "password": password})
    if resp.status_code == 302 and resp.url == "/login/otp/":
        code = re.search(r"\b(\d{6})\b", mail.outbox[-1].subject).group(1)
        resp = client.post("/login/otp/", {"code": code})
    return resp


@pytest.fixture
def as_user(client, users):
    def _as(name):
        client.force_login(users[name])
        return client

    return _as


def test_login_with_otp_and_role_landing(client, users, run_on_commit):
    resp = client.post("/login/", {"email": "l1@test.local", "password": PW})
    assert resp.url == "/login/otp/"
    assert client.post("/login/otp/", {"code": "000000"}).context["error"]
    code = re.search(r"\b(\d{6})\b", mail.outbox[-1].subject).group(1)
    resp = client.post("/login/otp/", {"code": code})
    assert resp.url == "/console/"
    client.logout()
    assert login(client, "vendora@test.local").url == "/vendor/"


def test_lockout_after_five_failures(client, users, run_on_commit):
    for _ in range(5):
        client.post("/login/", {"email": "l2@test.local", "password": "wrong"})
    resp = client.post("/login/", {"email": "l2@test.local", "password": PW})
    assert resp.status_code == 429


def test_password_reset_via_notify(client, users, run_on_commit):
    client.post("/password-reset/", {"email": "l3@test.local"})
    link = re.search(r"https?://[^\s\"]+/password-reset/[^\s\"]+", mail.outbox[-1].body).group(0)
    path = link.split("8000", 1)[-1] if "8000" in link else "/" + link.split("/", 3)[3]
    resp = client.get(path, follow=True)
    resp = client.post(resp.redirect_chain[-1][0], {"new_password1": "N3w-Passw0rd!x", "new_password2": "N3w-Passw0rd!x"})
    assert resp.url == "/password-reset/complete/"
    assert login(client, "l3@test.local", "N3w-Passw0rd!x").url == "/console/"


def test_idle_timeout_setting(settings):
    assert settings.SESSION_COOKIE_AGE == 1800 and settings.SESSION_SAVE_EVERY_REQUEST


@pytest.mark.parametrize("name", ["vm", "l1", "finance", "admin", "auditor"])
def test_console_pages_render(as_user, name, make_invoice):
    make_invoice(total=50000)
    c = as_user(name)
    from apps.vendors.models import Vendor

    for url in ["/console/", "/console/vendors/", "/console/invoices/", "/console/reports/ageing/", "/notifications/"] + [f"/console/vendors/{v.pk}/" for v in Vendor.objects.all()]:
        resp = c.get(url)
        assert resp.status_code in (200, 403), (url, resp.status_code)
    assert c.get("/console/").status_code == 200


def test_vendor_pages_render(as_user, make_invoice):
    inv = make_invoice(total=50000)
    c = as_user("vendora")
    for url in ["/vendor/", "/vendor/invoices/", "/vendor/payments/", "/vendor/profile/", "/vendor/bank/", "/vendor/users/", f"/vendor/invoices/{inv.pk}/", "/vendor/invoices/upload/"]:
        assert c.get(url).status_code == 200, url


def test_vendor_cannot_open_console(as_user):
    assert as_user("vendora").get("/console/").status_code == 404


def test_vendor_isolation_every_url(as_user, make_invoice, seed, users):
    inv_b = make_invoice(seed["vendor_b"], total=50000)
    payment_b = None
    for step, who in (("finance", "APPROVE"), ("l1", "APPROVE")):
        services.act(inv_b, by=users[step], action=who, fin_data={"tds_section": "NA"})
    payment_b = services.record_payment(inv_b, by=users["finance"], amount=Decimal("10000"), paid_on=inv_b.invoice_date, utr="U-B")
    c = as_user("vendora")
    urls = [
        f"/vendor/invoices/{inv_b.pk}/", f"/vendor/invoices/{inv_b.pk}/edit/", f"/vendor/invoices/{inv_b.pk}/parse-status/",
        f"/vendor/files/{inv_b.file_id}/", f"/console/invoices/{inv_b.pk}/", f"/console/files/{inv_b.file_id}/",
        f"/console/vendors/{seed['vendor_b'].pk}/",
    ]
    for url in urls:
        assert c.get(url).status_code == 404, url
    assert c.post(f"/vendor/invoices/{inv_b.pk}/discard/").status_code == 404
    listing = c.get("/vendor/invoices/").content.decode()
    assert inv_b.portal_ref not in listing
    export = c.get("/vendor/payments/?export=csv").content.decode()
    assert "U-B" not in export and payment_b.utr == "U-B"
    assert inv_b.portal_ref not in c.get("/vendor/invoices/?export=csv").content.decode()


def test_vendor_upload_flow_via_views(as_user, run_on_commit, seed):
    c = as_user("vendora")
    resp = c.post("/vendor/invoices/upload/", {"file": fake_pdf(invoice_no="VIEW-1", total="59000.00", taxable="50000.00", cgst="4500.00", sgst="4500.00")})
    assert resp.status_code == 302
    edit = c.get(resp.url)
    html = edit.content.decode()
    assert 'data-low-confidence="true"' in html and 'value="VIEW-1"' in html
    form = {"supplier_gstin": seed["vendor_a"].primary_gstin.gstin, "invoice_no": "VIEW-1", "invoice_date": edit.context["form"].initial["invoice_date"],
            "cost_centre": seed["cost_centre"].pk, "taxable": "50000.00", "cgst": "4500.00", "sgst": "4500.00", "igst": "0", "total": "59000.00"}
    resp = c.post(resp.url, form)
    assert resp.status_code == 302 and "/edit/" not in resp.url
    bad = c.post("/vendor/invoices/upload/", {"file": fake_pdf("x.exe")})
    assert bad.status_code == 302  # magic bytes are PDF; a non-PDF is rejected below
    from django.core.files.uploadedfile import SimpleUploadedFile

    resp = c.post("/vendor/invoices/upload/", {"file": SimpleUploadedFile("evil.pdf", b"MZ not a pdf")})
    assert "Only PDF, JPG and PNG" in resp.content.decode()


def test_exports_have_expected_columns(as_user, make_invoice):
    make_invoice(total=50000)
    c = as_user("finance")
    rows = list(csv.reader(io.StringIO(c.get("/console/invoices/?export=csv").content.decode())))
    assert rows[0][:5] == ["Portal ref", "Vendor code", "Vendor", "Supplier GSTIN", "Invoice no"] and len(rows) == 2
    wb = load_workbook(io.BytesIO(c.get("/console/reports/vendor-master/?export=xlsx").content))
    header = [cell.value for cell in next(wb.active.iter_rows(max_row=1))]
    assert header[:3] == ["Vendor code", "Legal name", "Trade name"]
    pend = list(csv.reader(io.StringIO(c.get("/console/reports/pending-by-approver/?export=csv").content.decode())))
    assert pend[0] == ["Approver", "Email", "Level", "Pending count", "Pending amount", "Oldest since"]
    assert any(r[1] == "finance2@test.local" for r in pend[1:])


def test_admin_toggles_template(as_user, run_on_commit, users):
    from apps.notify.models import EmailTemplate, Notification
    from apps.vendors import services as vs

    c = as_user("admin")
    c.post(reverse("console:admin_template_toggle", args=["vendor_blocked"]))
    assert not EmailTemplate.objects.get(code="vendor_blocked").enabled
    mail.outbox.clear()
    from apps.vendors.models import Vendor

    vs.block(Vendor.objects.get(vendor_code="V00001"), by=users["vm"], reason="x")
    assert not any("blocked" in m.subject for m in mail.outbox)
    assert Notification.objects.filter(template_code="vendor_blocked", status="SKIPPED").exists()
    c.post(reverse("console:admin_template_toggle", args=["login_otp"]), follow=True)
    assert EmailTemplate.objects.get(code="login_otp").enabled


def test_healthz(client, db):
    assert client.get("/healthz").json()["status"] == "ok"
